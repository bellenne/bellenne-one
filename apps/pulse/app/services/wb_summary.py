from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.clock import utc_now
from app.models import DailyMetric, Marketplace, MarketplaceAccount, WBSummaryMetric


MAX_WB_SUMMARY_BYTES = 15 * 1024 * 1024
MAX_WB_SUMMARY_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


@dataclass(frozen=True)
class WBSummaryRow:
    metric_date: date
    ordered_units: int
    ordered_amount: Decimal
    buyout_units: int
    buyout_amount: Decimal
    net_revenue: Decimal


@dataclass(frozen=True)
class WBSummaryParseResult:
    rows: list[WBSummaryRow]
    skipped_after_cutoff: int


@dataclass(frozen=True)
class WBSummaryImportResult:
    rows_imported: int
    first_date: date
    last_date: date
    skipped_after_cutoff: int


_EXPECTED_HEADERS = {
    "year": "год",
    "month": "месяц",
    "day": "день",
    "ordered_amount": (
        "сумма заказов по розничным ценам с учетом согласованной скидки руб"
    ),
    "ordered_units": "заказано шт",
    "buyout_amount": (
        "сумма продаж по розничным ценам с учетом согласованной скидки руб"
    ),
    "buyout_units": "выкупили шт",
    "net_revenue": "к перечислению за товар руб",
}


def _normalize_header(value: object) -> str:
    text = str(value or "").strip().casefold().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", " ", text).strip()


def _column_map(header: tuple[object, ...]) -> dict[str, int]:
    normalized = {_normalize_header(value): index for index, value in enumerate(header)}
    missing = [label for key, label in _EXPECTED_HEADERS.items() if label not in normalized]
    if missing:
        raise ValueError(
            "Файл не похож на «Сводный отчёт по продавцу» WB: "
            f"не найдены колонки {', '.join(missing)}."
        )
    return {key: normalized[label] for key, label in _EXPECTED_HEADERS.items()}


def _parse_excel_date(value: object, epoch) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        if isinstance(converted, datetime):
            return converted.date()
        return converted if isinstance(converted, date) else None
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
    return None


def _count(value: object, label: str, row_number: int) -> int:
    parsed = _number(value, label, row_number)
    if parsed != parsed.to_integral_value() or parsed < 0:
        raise ValueError(f"Строка {row_number}: некорректное значение «{label}».")
    return int(parsed)


def _number(value: object, label: str, row_number: int) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, bool):
        raise ValueError(f"Строка {row_number}: некорректное значение «{label}».")
    try:
        parsed = Decimal(str(value))
    except Exception as exc:
        raise ValueError(
            f"Строка {row_number}: некорректное значение «{label}»."
        ) from exc
    if not parsed.is_finite():
        raise ValueError(f"Строка {row_number}: некорректное значение «{label}».")
    return parsed


def parse_wb_summary_xlsx(content: bytes, through_date: date) -> WBSummaryParseResult:
    if not content:
        raise ValueError("Загружен пустой XLSX-файл.")
    if len(content) > MAX_WB_SUMMARY_BYTES:
        raise ValueError("XLSX-файл больше допустимых 15 МБ.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 1000 or sum(
                item.file_size for item in archive.infolist()
            ) > MAX_WB_SUMMARY_UNCOMPRESSED_BYTES:
                raise ValueError("Распакованный XLSX-файл слишком большой.")
    except BadZipFile as exc:
        raise ValueError("Не удалось прочитать XLSX-файл.") from exc
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise ValueError("Не удалось прочитать XLSX-файл.") from exc

    try:
        sheet = workbook.active
        # WB's generated workbook can declare the worksheet dimension as A1
        # even though the XML contains the full A:T table.  Read-only
        # openpyxl trusts that declaration unless dimensions are reset.
        if hasattr(sheet, "reset_dimensions"):
            sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("В XLSX-файле нет строк.")
        columns = _column_map(header)
        parsed_rows: list[WBSummaryRow] = []
        seen_dates: set[date] = set()
        skipped_after_cutoff = 0
        for row_number, row in enumerate(rows, start=2):
            day_value = row[columns["day"]] if columns["day"] < len(row) else None
            if day_value in (None, ""):
                # Year and month total rows are not daily facts.
                continue
            metric_date = _parse_excel_date(day_value, workbook.epoch)
            if metric_date is None:
                raise ValueError(f"Строка {row_number}: не удалось распознать дату.")
            try:
                report_year = int(row[columns["year"]])
                report_month = int(row[columns["month"]])
            except (TypeError, ValueError, IndexError) as exc:
                raise ValueError(f"Строка {row_number}: некорректные год или месяц.") from exc
            if (metric_date.year, metric_date.month) != (report_year, report_month):
                raise ValueError(
                    f"Строка {row_number}: дата не соответствует колонкам года и месяца."
                )
            if metric_date > through_date:
                skipped_after_cutoff += 1
                continue
            if metric_date in seen_dates:
                raise ValueError(
                    f"В файле несколько дневных строк за {metric_date:%d.%m.%Y}."
                )
            seen_dates.add(metric_date)
            parsed_rows.append(
                WBSummaryRow(
                    metric_date=metric_date,
                    ordered_units=_count(
                        row[columns["ordered_units"]],
                        _EXPECTED_HEADERS["ordered_units"],
                        row_number,
                    ),
                    ordered_amount=_number(
                        row[columns["ordered_amount"]],
                        _EXPECTED_HEADERS["ordered_amount"],
                        row_number,
                    ),
                    buyout_units=_count(
                        row[columns["buyout_units"]],
                        _EXPECTED_HEADERS["buyout_units"],
                        row_number,
                    ),
                    buyout_amount=_number(
                        row[columns["buyout_amount"]],
                        _EXPECTED_HEADERS["buyout_amount"],
                        row_number,
                    ),
                    net_revenue=_number(
                        row[columns["net_revenue"]],
                        _EXPECTED_HEADERS["net_revenue"],
                        row_number,
                    ),
                )
            )
    finally:
        workbook.close()

    if not parsed_rows:
        raise ValueError("В XLSX нет дневных строк до разрешённой даты.")
    parsed_rows.sort(key=lambda item: item.metric_date)
    return WBSummaryParseResult(parsed_rows, skipped_after_cutoff)


def apply_wb_summary_override(
    session: Session,
    account: MarketplaceAccount,
    metric_date: date,
) -> bool:
    summary = session.scalar(
        select(WBSummaryMetric).where(
            WBSummaryMetric.account_id == account.id,
            WBSummaryMetric.metric_date == metric_date,
        )
    )
    if summary is None:
        return False
    metric = session.scalar(
        select(DailyMetric).where(
            DailyMetric.account_id == account.id,
            DailyMetric.metric_date == metric_date,
            DailyMetric.scope_key == "total",
        )
    )
    if metric is None:
        metric = DailyMetric(
            user_id=account.user_id,
            account_id=account.id,
            marketplace=account.marketplace,
            metric_date=metric_date,
            scope_key="total",
        )
        session.add(metric)
    metric.ordered_units = summary.ordered_units
    metric.ordered_amount = summary.ordered_amount
    metric.buyout_units = summary.buyout_units
    metric.buyout_amount = summary.buyout_amount
    metric.net_revenue = summary.net_revenue
    metric.source = (
        "api+wb_summary"
        if metric.source and "api" in metric.source.split("+")
        else "wb_summary"
    )
    metric.collected_at = utc_now()
    return True


def import_wb_summary(
    session: Session,
    account: MarketplaceAccount,
    content: bytes,
    source_name: str,
    through_date: date,
) -> WBSummaryImportResult:
    if account.marketplace != Marketplace.WB.value or account.is_demo:
        raise ValueError("Сводный отчёт можно загрузить только в рабочий кабинет WB.")
    parsed = parse_wb_summary_xlsx(content, through_date)
    safe_source_name = (source_name.strip() or "WB summary.xlsx")[:255]
    for item in parsed.rows:
        summary = session.scalar(
            select(WBSummaryMetric).where(
                WBSummaryMetric.account_id == account.id,
                WBSummaryMetric.metric_date == item.metric_date,
            )
        )
        if summary is None:
            summary = WBSummaryMetric(
                user_id=account.user_id,
                account_id=account.id,
                metric_date=item.metric_date,
                source_name=safe_source_name,
            )
            session.add(summary)
        summary.ordered_units = item.ordered_units
        summary.ordered_amount = item.ordered_amount
        summary.buyout_units = item.buyout_units
        summary.buyout_amount = item.buyout_amount
        summary.net_revenue = item.net_revenue
        summary.source_name = safe_source_name
        summary.imported_at = utc_now()
        session.flush()
        apply_wb_summary_override(session, account, item.metric_date)

    return WBSummaryImportResult(
        rows_imported=len(parsed.rows),
        first_date=parsed.rows[0].metric_date,
        last_date=parsed.rows[-1].metric_date,
        skipped_after_cutoff=parsed.skipped_after_cutoff,
    )
